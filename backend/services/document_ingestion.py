"""Bounded deterministic extraction. Never executes macros or student code."""
import io
import re
import zipfile
from pathlib import Path
import pymupdf
from docx import Document
from pptx import Presentation

MAX_BYTES = 30 * 1024 * 1024
MAX_PAGES = 50
MAX_EXPANDED_BYTES = 60 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
CODE_SUFFIXES = {"py", "js", "ts", "tsx", "jsx", "java", "c", "cpp", "h", "rs", "go", "sql", "txt"}


def validate_ooxml(content: bytes, kind: str):
    if not content.startswith(b"PK"):
        raise ValueError("Not a valid Office document.")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 3000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
            raise ValueError("Office archive exceeds processing limits.")
        names = {entry.filename for entry in entries}
        expected = "word/document.xml" if kind == "docx" else "ppt/presentation.xml"
        if expected not in names or "[Content_Types].xml" not in names:
            raise ValueError("Office document type does not match its contents.")
        for entry in entries:
            parts = Path(entry.filename).parts
            if entry.filename.startswith(("/", "\\")) or ".." in parts or "\\" in entry.filename or entry.flag_bits & 1:
                raise ValueError("Invalid Office archive entry.")
            name = entry.filename.casefold()
            if any(marker in name for marker in ("vbaproject", "/embeddings/", "activex")):
                raise ValueError("Macros and embedded executable objects are not supported.")
            if name.endswith(".xml") or name.endswith(".rels"):
                if entry.file_size > 5 * 1024 * 1024:
                    raise ValueError("Office XML part exceeds processing limits.")
                data = archive.read(entry)
                if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                    raise ValueError("XML entity declarations are not supported.")


def ingest(content: bytes, filename: str, allow_code: bool = False) -> dict:
    if not content:
        raise ValueError("The uploaded document is empty.")
    if len(content) > MAX_BYTES:
        raise ValueError("Document size must not exceed 30 MB.")
    kind = Path(filename).suffix.lower().lstrip(".")
    if allow_code and kind in CODE_SUFFIXES | {"zip"}:
        return ingest_code(content, filename)
    if kind not in {"pdf", "docx", "pptx"}:
        raise ValueError("Upload a PDF, DOCX or PPTX document.")
    pages = []
    warnings = []
    if kind == "pdf":
        if not content.lstrip().startswith(b"%PDF-"):
            raise ValueError("The uploaded file is not a PDF.")
        with pymupdf.open(stream=content, filetype="pdf") as document:
            if document.needs_pass or not 1 <= document.page_count <= MAX_PAGES:
                raise ValueError("PDF must be unencrypted and contain 1–50 pages.")
            total = 0
            for index, page in enumerate(document):
                text = page.get_text()
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise ValueError("Document text exceeds extraction limits.")
                blocks = [{"text": block[4], "bbox": list(block[:4])} for block in page.get_text("blocks") if len(block) > 6 and block[6] == 0]
                pages.append({"number": index + 1, "unit": "page", "text": text, "rotation": page.rotation, "width": page.rect.width, "height": page.rect.height, "regions": blocks})
                if not text.strip():
                    warnings.append(f"Page {index + 1}: no embedded text; handwriting/scan interpretation is not available in deterministic extraction.")
    else:
        validate_ooxml(content, kind)
        if kind == "docx":
            document = Document(io.BytesIO(content))
            # DOCX has no reliable page boundaries without a layout renderer.
            paragraphs = [{"text": paragraph.text, "style": paragraph.style.name if paragraph.style else ""} for paragraph in document.paragraphs]
            tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
            text = "\n".join([row["text"] for row in paragraphs] + ["\n".join("\t".join(row) for row in table) for table in tables])
            pages = [{"number": 1, "unit": "document", "text": text, "paragraphs": paragraphs, "tables": tables, "regions": []}]
            warnings.append("DOCX page numbers and visual layout are unavailable; references use document paragraphs/tables.")
        else:
            presentation = Presentation(io.BytesIO(content))
            if not 1 <= len(presentation.slides) <= MAX_PAGES:
                raise ValueError("PPTX must contain 1–50 slides.")
            for index, slide in enumerate(presentation.slides):
                texts = []
                tables = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text)
                    if shape.has_table:
                        tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
                text = "\n".join(texts + ["\n".join("\t".join(row) for row in table) for table in tables])
                pages.append({"number": index + 1, "unit": "slide", "text": text, "tables": tables, "regions": []})
            warnings.append("PPTX text and tables extracted. Slide rendering, visuals, speaking quality and delivery are not evaluated.")
    text = "\n".join(page["text"] for page in pages)
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Document text exceeds extraction limits.")
    for page in pages:
        page["headings"] = [line.strip() for line in page["text"].splitlines() if re.match(r"^(?:Question\s+\d+|\d+[.)]\s|Abstract\b|Introduction\b|Discussion\b|Conclusions?\b|References\b)", line.strip(), re.I)][:100]
    return {"version": "deterministic-v1", "file_type": kind, "unit_count": len(pages), "pages": pages, "extracted_text": text, "warnings": warnings, "question_associations": [], "requires_human_review": True}


def ingest_code(content: bytes, filename: str) -> dict:
    kind = Path(filename).suffix.lower().lstrip(".")
    files = []
    if kind == "zip":
        if not content.startswith(b"PK"):
            raise ValueError("Invalid source archive.")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(entry.file_size for entry in entries) > 5 * 1024 * 1024:
                raise ValueError("Source archive exceeds 200 files or 5 MB expanded.")
            for entry in entries:
                path = Path(entry.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in entry.filename or entry.flag_bits & 1 or (entry.external_attr >> 16) & 0o170000 == 0o120000:
                    raise ValueError("Unsafe source archive entry.")
                if entry.is_dir():
                    continue
                if path.suffix.lower().lstrip(".") not in CODE_SUFFIXES:
                    raise ValueError("Archive contains an unsupported source file.")
                files.append((entry.filename, archive.read(entry)))
    else:
        files = [(Path(filename).name, content)]
    pages = []
    total = 0
    for name, data in files:
        text = data.decode("utf-8")
        if "\x00" in text:
            raise ValueError("Source files must contain UTF-8 text.")
        total += len(text)
        if total > MAX_TEXT_CHARS:
            raise ValueError("Source text exceeds extraction limits.")
        pages.append({"number": len(pages) + 1, "unit": "file", "file_name": name, "text": text, "regions": [], "headings": []})
    if not pages:
        raise ValueError("Source archive is empty.")
    return {"version": "deterministic-v1", "file_type": kind, "unit_count": len(pages), "pages": pages,
            "extracted_text": "\n".join(page["text"] for page in pages), "question_associations": [],
            "warnings": ["Source files are available for static review only. Code is never executed; runtime correctness is not evaluated."], "requires_human_review": True}
