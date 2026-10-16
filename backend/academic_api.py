import io
import re
import copy
import math
import logging
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Literal
from uuid import uuid4, UUID

import pymupdf
from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, Query, Request, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from persistence import (
    AcademicAssignment,
    AcademicCourse,
    AcademicGrade,
    AcademicSubmission,
    Answer,
    Assignment,
    AssignmentFamily,
    CourseEnrollment,
    CourseOffering,
    OriginalityReport,
    OriginalityAnalysis,
    SessionLocal,
    Submission,
    blob_exists,
    read_blob,
    now_utc,
    save_blob,
    write_blob,
)
from assignment_api import DEMO_MODE
from auth import auth_required, request_user
from account_models import CourseOwner, OfferingPolicy, EnrollmentState
from permissions import can_teach, enrollment_active
from courses import new_code, join_code
from services.document_ingestion import ingest, CODE_SUFFIXES
from document_models import SubmissionDocument, PublishedDocumentReview
from document_review import published_feedback
from student_identity import student_display_name
from starlette.concurrency import run_in_threadpool
import hashlib


router = APIRouter(
    prefix="/academic",
    tags=["Academic"],
)

DEMO_STUDENT_ID = "demo-student-001"
DEMO_STUDENT_NAME = "Student 01"
COURSE_ID = "phy101-demo-course"
OFFERING_IDS = {
    "2024": "phy101-fall-2024",
    "2025": "phy101-fall-2025",
    "2026": "phy101-fall-2026",
}
FAMILY_ID = "phy101-force-family"
ASSIGNMENT_IDS = {
    "2024": "phy101-force-assignment-2024",
    "2025": "phy101-force-assignment-2025",
    "2026": "phy101-force-assignment-2026",
}
SEED_SUBMISSION_IDS = {
    "2024": "phy101-history-submission-2024",
    "2025": "phy101-history-submission-2025",
}
MAX_UPLOAD_BYTES = 30 * 1024 * 1024
MIN_MATCH_TOKENS = 8
TOKEN_RE = re.compile(r"[^\W_]+")
ALLOWED_REVIEW_STATUSES = {"unreviewed", "reviewed", "follow_up"}


class OfferingCreate(BaseModel):
    course_id: str = Field(min_length=1, max_length=36)
    academic_year: str = Field(min_length=1, max_length=12)
    term: str = Field(min_length=1, max_length=32)
    section: str = Field(default="01", min_length=1, max_length=24)
    status: Literal["active", "archived", "upcoming"] = "active"
    class_code: str | None = Field(default=None, max_length=32)


class EnrollmentRequest(BaseModel):
    class_code: str = Field(min_length=3, max_length=32)


class CriterionPayload(BaseModel):
    id: str | None = Field(default=None, max_length=80)
    criterion_id: str | None = Field(default=None, max_length=80)
    description: str = Field(min_length=1, max_length=1000)
    maximum_points: float = Field(gt=0, le=10000, allow_inf_nan=False)
    evidence_expected: str = Field(default="", max_length=2000)
    partial_credit: str = Field(default="", max_length=2000)


class QuestionPayload(BaseModel):
    id: str | None = Field(default=None, max_length=80)
    title: str = Field(min_length=1, max_length=200)
    label: str | None = Field(default=None, max_length=200)
    statement: str = Field(default="", max_length=10000)
    reference_solution: str = Field(default="", max_length=20000)
    maximum_points: float = Field(gt=0, le=10000, allow_inf_nan=False)
    criteria: list[CriterionPayload] = Field(default_factory=list, max_length=50)
    region: dict | None = None


class AcademicAssignmentCreate(BaseModel):
    offering_id: str = Field(min_length=1, max_length=36)
    title: str = Field(min_length=1, max_length=200)
    assignment_type: Literal["homework", "problem_set", "essay", "lab_report", "presentation", "research_report", "programming", "general"] = "homework"
    description: str = Field(default="", max_length=5000)
    available_at: datetime | None = None
    due_at: datetime | None = None
    maximum_score: float = Field(default=100, gt=0, le=10000, allow_inf_nan=False)
    allow_resubmission: bool = False
    originality_enabled: bool = True
    originality_scope: Literal["family_then_course", "family", "course"] = "family_then_course"
    questions: list[QuestionPayload] = Field(default_factory=list, max_length=20)


class AssignmentReuse(BaseModel):
    offering_id: str = Field(min_length=1, max_length=36)
    available_at: datetime | None = None
    due_at: datetime | None = None
    copy_rubric: bool = True
    copy_reference_solution: bool = True
    copy_template: bool = True


class ReviewUpdate(BaseModel):
    review_status: Literal["unreviewed", "reviewed", "follow_up"]
    teacher_notes: str = Field(default="", max_length=10000)


class CourseOfferingResponse(BaseModel):
    id: str
    academic_year: str
    term: str
    section: str
    status: str
    pending_assignments: int
    enrollment_count: int | None = None


class CourseItemResponse(BaseModel):
    id: str
    code: str
    name: str
    instructor: str
    department: str
    offerings: list[CourseOfferingResponse]


class CoursesResponse(BaseModel):
    items: list[CourseItemResponse]


class CourseIdentityResponse(BaseModel):
    id: str
    code: str
    name: str
    instructor: str
    department: str


class AcademicAssignmentResponse(BaseModel):
    id: str
    offering_id: str
    family_id: str
    title: str
    assignment_type: str
    description: str
    available_at: str | None
    due_at: str | None
    maximum_score: float
    allow_resubmission: bool
    originality_enabled: bool
    originality_scope: str
    questions: list[dict]
    template_regions: list[dict]
    template_rotation: int
    submission_count: int
    created_at: str | None
    reused_from: str | None = None
    copied: dict | None = None
    legacy_assignment_id: str | None = None


class AssignmentItemsResponse(BaseModel):
    items: list[AcademicAssignmentResponse]


class OfferingDetailResponse(BaseModel):
    enrollment_count: int | None = None
    id: str
    academic_year: str
    term: str
    section: str
    status: str
    class_code: str
    course: CourseIdentityResponse
    instructor: str
    assignments: list[AcademicAssignmentResponse]


class OriginalityEvidenceResponse(BaseModel):
    matched_tokens: int
    source_offsets: dict[str, int]
    candidate_offsets: dict[str, int]
    source_excerpt: str
    candidate_excerpt: str


class OriginalityReportResponse(BaseModel):
    analysis_status: str
    id: str
    submission_id: str
    candidate_submission_id: str
    student_id: str | None
    candidate_student_id: str | None
    assignment_family_id: str
    match_scope: str
    match_ratio: float | None
    eligible_token_count: int
    matched_token_count: int
    ratio_definition: str
    offset_unit: str
    minimum_passage_tokens: int
    evidence: list[OriginalityEvidenceResponse]
    review_status: str
    teacher_notes: str
    created_at: str | None
    reviewed_at: str | None
    interpretation: str
    extracted_text: str
    document_title: str
    assignment_title: str
    course_name: str
    course_code: str
    academic_year: str
    term: str
    source_title: str
    source_assignment: str
    source_course: str
    source_term: str


class OriginalityItemsResponse(BaseModel):
    items: list[OriginalityReportResponse]


class SubmissionResponse(BaseModel):
    id: str
    assignment_id: str
    student_id: str
    file_name: str
    file_type: str
    status: str
    extracted_text: str
    submitted_at: str | None
    originality_reports: list[OriginalityReportResponse]


class AcademicStudentGradeResponse(BaseModel):
    id: str
    assignment_id: str
    assignment_title: str
    submission_id: str
    points: float
    maximum_score: float
    published_at: str | None
    criteria: list = Field(default_factory=list)
    feedback: str = ""
    rubric: list = Field(default_factory=list)


class StudentGradeItemsResponse(BaseModel):
    items: list[AcademicStudentGradeResponse]


class StudentAssignmentResponse(AcademicAssignmentResponse):
    submission: dict | None
    published_grade: dict | None


class StudentAssignmentItemsResponse(BaseModel):
    items: list[StudentAssignmentResponse]


class OriginalityAnalysisResponse(BaseModel):
    analysis_status: Literal['missing', 'queued', 'analyzing', 'failed', 'insufficient_text', 'no_sources', 'disabled', 'incomplete', 'completed']
    submission_id: str
    document_title: str
    assignment_title: str
    course_name: str
    academic_year: str
    extracted_text: str
    eligible_token_count: int | None
    candidate_count: int | None
    match_ratio: float | None
    match_scope: str
    created_at: str | None


class SubmissionOriginalityResponse(BaseModel):
    submission_id: str
    items: list[OriginalityReportResponse]
    analysis: OriginalityAnalysisResponse


class PublishResponse(BaseModel):
    assignment_id: str
    published_count: int
    grades: list[dict]
    notice: str


class EnrollmentResponse(BaseModel):
    id: str
    offering_id: str
    student_id: str
    student_name: str
    created: bool
    status: str = "active"


class OfferingCreateResponse(BaseModel):
    id: str
    course_id: str
    academic_year: str
    term: str
    section: str
    status: str
    class_code: str


class ArchiveSubmissionResponse(BaseModel):
    student_name: str
    submission_id: str
    assignment_id: str
    assignment_title: str
    academic_year: str
    term: str
    student_id: str
    file_name: str
    file_type: str
    submitted_at: str | None
    extracted_text: str


class ArchiveItemsResponse(BaseModel):
    items: list[ArchiveSubmissionResponse]


def _as_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _timestamp(value: datetime | None) -> str | None:
    return _as_utc(value).isoformat() if value is not None else None


def _seed_questions() -> list[dict]:
    return [
        {
            "id": "q-force",
            "title": "Resolve the forces",
            "label": "Force components",
            "statement": "A block slides down a frictionless 30 degree incline. Resolve its weight.",
            "reference_solution": "The component parallel to the incline is mg sin(theta); the normal component is mg cos(theta).",
            "maximum_points": 10,
            "criteria": [
                {"id": "parallel", "criterion_id": "parallel", "description": "Correct parallel component", "maximum_points": 5},
                {"id": "normal", "criterion_id": "normal", "description": "Correct normal component", "maximum_points": 5},
            ],
            "region": None,
        }
    ]


def _seed_pdf(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=612, height=792)
    page.insert_textbox(pymupdf.Rect(48, 48, 564, 740), text, fontsize=11)
    content = document.tobytes()
    document.close()
    return content


def _seed_docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _extract_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    blocks.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(text for text in blocks if text.strip())


def ensure_academic_seed() -> None:
    if not DEMO_MODE:
        return

    with SessionLocal() as db:
        if db.get(AcademicCourse, COURSE_ID) is None:
            db.add(
                AcademicCourse(
                    id=COURSE_ID,
                    code="PHY101",
                    name="General Physics I",
                    instructor="Dr. Ada Rivera",
                    department="Physics",
                )
            )
        if db.get(AssignmentFamily, FAMILY_ID) is None:
            db.add(
                AssignmentFamily(
                    id=FAMILY_ID,
                    course_id=COURSE_ID,
                    title="Forces on an Incline",
                )
            )
        db.flush()
        for year, offering_id in OFFERING_IDS.items():
            if db.get(CourseOffering, offering_id) is None:
                db.add(
                    CourseOffering(
                        id=offering_id,
                        course_id=COURSE_ID,
                        academic_year=year,
                        term="Fall",
                        section="01",
                        status="archived" if year != "2026" else "active",
                        class_code=f"PHY101-F{year[-2:]}",
                    )
                )
        db.flush()

        old_docx_fixture = db.get(AcademicSubmission, SEED_SUBMISSION_IDS["2025"])
        if old_docx_fixture and old_docx_fixture.file_type != "docx":
            relative = Path(ASSIGNMENT_IDS["2025"]) / "academic" / "seed-2025.docx"
            write_blob(relative.as_posix(), _seed_docx(old_docx_fixture.extracted_text))
            old_docx_fixture.file_path = relative.as_posix()
            old_docx_fixture.file_name = "seed-2025.docx"
            old_docx_fixture.file_type = "docx"
            old_docx_fixture.extracted_text = _extract_docx(
                read_blob(relative.as_posix())
            )
            legacy = db.get(Submission, old_docx_fixture.legacy_submission_id)
            if legacy:
                legacy.source_path = ""

        seed_texts = {
            "2024": (
                "Forces on an Incline\n"
                "For a block on a frictionless incline, choose axes parallel and perpendicular to the surface. "
                "The gravitational force points vertically downward, so resolve it into components along these axes. "
                "The parallel component is mg sin(theta), directed down the slope. "
                "The perpendicular component is mg cos(theta), directed into the surface. "
                "Because there is no acceleration perpendicular to the plane, the normal force balances that component. "
                "Newton's second law along the plane then gives acceleration g sin(theta)."
            ),
            "2025": (
                "A short worked example uses the shared formula F = ma. "
                "For a cart moving at constant speed, the net force is zero."
            ),
        }
        legacy_ids: dict[str, str] = {}
        for year, assignment_id in ASSIGNMENT_IDS.items():
            legacy_id = f"legacy-{assignment_id}"
            legacy_ids[year] = legacy_id
            if db.get(Assignment, legacy_id) is None:
                db.add(
                    Assignment(
                        id=legacy_id,
                        title="Forces on an Incline",
                        subject="Physics",
                        description="PHY101 reusable assignment family.",
                        status="active" if year == "2026" else "completed",
                        questions=_seed_questions(),
                        template_regions=[],
                        template_rotation=0,
                    )
                )
        db.flush()

        for year, assignment_id in ASSIGNMENT_IDS.items():
            legacy_id = legacy_ids[year]
            academic_assignment = db.get(AcademicAssignment, assignment_id)
            if academic_assignment is None:
                db.add(
                    AcademicAssignment(
                        id=assignment_id,
                        offering_id=OFFERING_IDS[year],
                        family_id=FAMILY_ID,
                        legacy_assignment_id=legacy_id,
                        title="Forces on an Incline",
                        assignment_type="problem_set",
                        description="Resolve gravity into components on a frictionless incline.",
                        available_at=None,
                        due_at=None,
                        maximum_score=10,
                        allow_resubmission=False,
                        originality_enabled=True,
                        originality_scope="family_then_course",
                        questions=_seed_questions(),
                        template_regions=[],
                        template_rotation=0,
                    )
                )
            elif academic_assignment.legacy_assignment_id is None:
                academic_assignment.legacy_assignment_id = legacy_id
        db.flush()

        extracted_texts: dict[str, str] = {}
        legacy_submission_ids: dict[str, str] = {}
        for year, submission_id in SEED_SUBMISSION_IDS.items():
            if db.get(AcademicSubmission, submission_id) is not None:
                continue
            assignment_id = ASSIGNMENT_IDS[year]
            text = seed_texts[year]
            # New synthetic IDs must fit submissions.id's PostgreSQL VARCHAR(36).
            # Existing academic fixtures above retain their original references.
            legacy_submission_id = f"legacy-phy101-submission-{year}"
            legacy_submission_ids[year] = legacy_submission_id
            file_type = "pdf" if year == "2024" else "docx"
            relative = Path(assignment_id) / "academic" / f"seed-{year}.{file_type}"
            if not blob_exists(relative.as_posix()):
                write_blob(
                    relative.as_posix(),
                    _seed_pdf(text) if file_type == "pdf" else _seed_docx(text)
                )
            if file_type == "pdf":
                document = pymupdf.open(
                    stream=read_blob(relative.as_posix()), filetype="pdf"
                )
                extracted_text = "\n".join(page.get_text() for page in document)
                document.close()
            else:
                extracted_text = _extract_docx(read_blob(relative.as_posix()))
            extracted_texts[year] = extracted_text
            text_path = Path(assignment_id, "academic", f"seed-{year}.txt").as_posix()
            if not blob_exists(text_path):
                write_blob(text_path, extracted_text.encode("utf-8"))

            if db.get(Submission, legacy_submission_id) is None:
                db.add(
                    Submission(
                        id=legacy_submission_id,
                        assignment_id=legacy_ids[year],
                        student_name=f"Historical Student {year}",
                        anonymous_student_id=f"Historical Student {year}",
                        source_path=relative.as_posix() if file_type == "pdf" else "",
                        page_count=1,
                        status="ready",
                    )
                )
        db.flush()

        for year, submission_id in SEED_SUBMISSION_IDS.items():
            if db.get(AcademicSubmission, submission_id) is not None:
                continue
            assignment_id = ASSIGNMENT_IDS[year]
            file_type = "pdf" if year == "2024" else "docx"
            relative = Path(assignment_id) / "academic" / f"seed-{year}.{file_type}"
            db.add(
                AcademicSubmission(
                    id=submission_id,
                    assignment_id=assignment_id,
                    student_id=f"historical-student-{year}",
                    student_name=f"Historical Student {year}",
                    legacy_submission_id=legacy_submission_ids[year],
                    file_path=relative.as_posix(),
                    file_name=f"seed-{year}.{file_type}",
                    file_type=file_type,
                    extracted_text=extracted_texts[year],
                    status="submitted",
                )
            )
        db.flush()
        for year in OFFERING_IDS:
            enrollment_exists = db.scalar(
                select(CourseEnrollment.id).where(
                    CourseEnrollment.offering_id == OFFERING_IDS[year],
                    CourseEnrollment.student_id == DEMO_STUDENT_ID,
                )
            )
            if enrollment_exists is None:
                db.add(
                    CourseEnrollment(
                        offering_id=OFFERING_IDS[year],
                        student_id=DEMO_STUDENT_ID,
                        student_name=DEMO_STUDENT_NAME,
                    )
                )
        db.commit()


def _student_identity(request: Request) -> tuple[str, str]:
    if auth_required():
        user = request_user(request)
        if user.role != "student" or user.student_id is None:
            raise HTTPException(status_code=403, detail="Student role required.")
        from student_identity import student_display_name
        with SessionLocal() as db:
            return user.student_id, student_display_name(db, user.student_id)
    return DEMO_STUDENT_ID, DEMO_STUDENT_NAME


def _student_view(assignment: dict) -> dict:
    result = copy.deepcopy(assignment)
    for question in result.get("questions", []):
        if isinstance(question, dict):
            question.pop("reference_solution", None)
    return result


def _active_enrollment_count(db, offering_id):
    rows = db.scalars(select(CourseEnrollment).where(CourseEnrollment.offering_id == offering_id)).all()
    return len({row.student_id for row in rows if enrollment_active(db, row)})


def _course_dict(db: Session, course: AcademicCourse, role: str, student_id: str) -> dict:
    offerings = db.scalars(
        select(CourseOffering)
        .where(CourseOffering.course_id == course.id)
        .order_by(CourseOffering.academic_year)
    ).all()
    items = []
    for offering in offerings:
        if role == "student" and not enrollment_active(db, db.scalar(
            select(CourseEnrollment).where(
                CourseEnrollment.offering_id == offering.id,
                CourseEnrollment.student_id == student_id,
            )
        )):
            continue
        assignments = db.scalars(
            select(AcademicAssignment).where(AcademicAssignment.offering_id == offering.id)
        ).all()
        submitted = {
            row.assignment_id
            for row in db.scalars(
                select(AcademicSubmission).where(AcademicSubmission.student_id == student_id)
            ).all()
        }
        pending = sum(assignment.id not in submitted for assignment in assignments)
        items.append(
            CourseOfferingResponse(
                id=offering.id,
                academic_year=offering.academic_year,
                term=offering.term,
                section=offering.section,
                status=offering.status,
                pending_assignments=pending,
                enrollment_count=_active_enrollment_count(db, offering.id) if role == "teacher" else None,
            ).model_dump()
        )
    return {
        "id": course.id,
        "code": (db.get(CourseOwner, course.id).public_code if db.get(CourseOwner, course.id) else course.code),
        "name": course.name,
        "instructor": course.instructor,
        "department": course.department,
        "offerings": items,
    }


def _assignment_view(db: Session, assignment: AcademicAssignment) -> dict:
    submissions = db.scalars(
        select(AcademicSubmission).where(AcademicSubmission.assignment_id == assignment.id)
    ).all()
    return {
        "id": assignment.id,
        "offering_id": assignment.offering_id,
        "family_id": assignment.family_id,
        "title": assignment.title,
        "assignment_type": assignment.assignment_type,
        "description": assignment.description,
        "available_at": _timestamp(assignment.available_at),
        "due_at": _timestamp(assignment.due_at),
        "maximum_score": assignment.maximum_score,
        "allow_resubmission": assignment.allow_resubmission,
        "originality_enabled": assignment.originality_enabled,
        "originality_scope": assignment.originality_scope,
        "questions": assignment.questions or [],
        "template_regions": assignment.template_regions or [],
        "template_rotation": assignment.template_rotation,
        "submission_count": len(submissions),
        "created_at": _timestamp(assignment.created_at),
        "legacy_assignment_id": assignment.legacy_assignment_id,
    }


def _normalize_questions(
    questions: list[QuestionPayload], maximum_score: float
) -> list[dict]:
    if not questions:
        questions = [
            QuestionPayload(
                id="q1",
                title="Ответ на задание",
                statement="Выполните задание.",
                maximum_points=maximum_score,
            )
        ]
    normalized = []
    ids: set[str] = set()
    for index, question in enumerate(questions):
        item = question.model_dump()
        item["id"] = item["id"] or f"q{index + 1}"
        if item["id"] in ids:
            raise HTTPException(status_code=422, detail="Question IDs must be unique.")
        ids.add(item["id"])
        item["label"] = item["label"] or item["title"]
        criteria = []
        criterion_ids = set()
        for criterion_index, criterion in enumerate(item["criteria"]):
            criterion_id = (
                criterion["criterion_id"]
                or criterion["id"]
                or f"{item['id']}-c{criterion_index + 1}"
            )
            if criterion_id in criterion_ids:
                raise HTTPException(422, "Criterion IDs must be unique within a question.")
            criterion_ids.add(criterion_id)
            criteria.append(
                {
                    "id": criterion_id,
                    "criterion_id": criterion_id,
                    "description": criterion["description"],
                    "maximum_points": criterion["maximum_points"],
                    "evidence_expected": criterion.get("evidence_expected", ""),
                    "partial_credit": criterion.get("partial_credit", ""),
                }
            )
        if criteria and sum(row["maximum_points"] for row in criteria) > item["maximum_points"]:
            raise HTTPException(status_code=422, detail="Rubric criteria exceed the question maximum.")
        if not criteria:
            criteria = [
                {
                    "id": f"{item['id']}-overall",
                    "criterion_id": f"{item['id']}-overall",
                    "description": "Общая корректность",
                    "maximum_points": item["maximum_points"],
                }
            ]
        item["criteria"] = criteria
        normalized.append(item)
    if sum(question["maximum_points"] for question in normalized) > maximum_score + 1e-6:
        raise HTTPException(status_code=422, detail="Question maximums exceed assignment maximum_score.")
    return normalized


def _new_academic_assignment(
    db: Session,
    offering: CourseOffering,
    title: str,
    assignment_type: str,
    description: str,
    available_at: datetime | None,
    due_at: datetime | None,
    maximum_score: float,
    allow_resubmission: bool,
    originality_enabled: bool,
    originality_scope: str,
    questions: list[dict],
    template_regions: list | None = None,
    template_rotation: int = 0,
    family_id: str | None = None,
) -> AcademicAssignment:
    if available_at and due_at and _as_utc(available_at) > _as_utc(due_at):
        raise HTTPException(status_code=422, detail="available_at must not be after due_at.")
    family = db.get(AssignmentFamily, family_id) if family_id else None
    if family_id and family is None:
        raise HTTPException(status_code=404, detail="Assignment family not found.")
    if family and family.course_id != offering.course_id:
        raise HTTPException(status_code=422, detail="Assignment family belongs to another course.")
    family = family or AssignmentFamily(
        course_id=offering.course_id,
        title=title,
    )
    db.add(family)
    db.flush()
    legacy = Assignment(
        title=title,
        subject=(db.get(AcademicCourse, offering.course_id).department or "General"),
        description=description,
        status="active",
        questions=questions,
        template_regions=template_regions or [],
        template_rotation=template_rotation,
    )
    db.add(legacy)
    db.flush()
    result = AcademicAssignment(
        offering_id=offering.id,
        family_id=family.id,
        legacy_assignment_id=legacy.id,
        title=title,
        assignment_type=assignment_type,
        description=description,
        available_at=_as_utc(available_at),
        due_at=_as_utc(due_at),
        maximum_score=maximum_score,
        allow_resubmission=allow_resubmission,
        originality_enabled=originality_enabled,
        originality_scope=originality_scope,
        questions=questions,
        template_regions=template_regions or [],
        template_rotation=template_rotation,
    )
    db.add(result)
    db.flush()
    return result


def _tokens(text: str) -> tuple[list[str], list[tuple[int, int]]]:
    matches = list(TOKEN_RE.finditer(text))
    return [match.group(0).casefold() for match in matches], [(match.start(), match.end()) for match in matches]


def _match_evidence(source_text: str, candidate_text: str) -> tuple[int, int, list[dict]]:
    source_tokens, source_offsets = _tokens(source_text)
    candidate_tokens, candidate_offsets = _tokens(candidate_text)
    eligible = len(source_tokens)
    if eligible < MIN_MATCH_TOKENS or len(candidate_tokens) < MIN_MATCH_TOKENS:
        return 0, eligible, []
    matcher = SequenceMatcher(None, source_tokens, candidate_tokens, autojunk=False)
    blocks = [block for block in matcher.get_matching_blocks() if block.size >= MIN_MATCH_TOKENS]
    evidence = []
    matched = 0
    last_source_end = -1
    for block in blocks:
        source_start = source_offsets[block.a][0]
        source_end = source_offsets[block.a + block.size - 1][1]
        candidate_start = candidate_offsets[block.b][0]
        candidate_end = candidate_offsets[block.b + block.size - 1][1]
        if block.a < last_source_end:
            continue
        last_source_end = block.a + block.size
        matched += block.size
        evidence.append(
            {
                "matched_tokens": block.size,
                "source_offsets": {"start": source_start, "end": source_end},
                "candidate_offsets": {"start": candidate_start, "end": candidate_end},
                "source_excerpt": source_text[source_start:source_end],
                "candidate_excerpt": candidate_text[candidate_start:candidate_end],
            }
        )
    return min(matched, eligible), eligible, evidence


def _compare_submission_pairs(
    db: Session, submission: AcademicSubmission
) -> list[OriginalityReport]:
    if not submission.extracted_text.strip():
        return []
    assignment = db.get(AcademicAssignment, submission.assignment_id)
    if assignment is None or not assignment.originality_enabled:
        return []
    offering = db.get(CourseOffering, assignment.offering_id)
    family = db.get(AssignmentFamily, assignment.family_id)
    candidate_rows = db.scalars(
        select(AcademicSubmission)
        .join(AcademicAssignment, AcademicSubmission.assignment_id == AcademicAssignment.id)
        .join(CourseOffering, AcademicAssignment.offering_id == CourseOffering.id)
        .where(
            CourseOffering.course_id == offering.course_id,
            AcademicSubmission.id != submission.id,
            AcademicSubmission.student_id != submission.student_id,
        )
        .order_by(AcademicSubmission.created_at, AcademicSubmission.id)
    ).all()
    family_matches: list[tuple[AcademicSubmission, int, int, list[dict]]] = []
    course_matches: list[tuple[AcademicSubmission, int, int, list[dict]]] = []
    compared_family = compared_course = 0
    for candidate in candidate_rows:
        candidate_assignment = db.get(AcademicAssignment, candidate.assignment_id)
        if candidate_assignment is None:
            continue
        if len(_tokens(candidate.extracted_text)[0]) < MIN_MATCH_TOKENS:
            continue
        if candidate_assignment.family_id == family.id:
            compared_family += 1
        else:
            compared_course += 1
        match = _match_evidence(submission.extracted_text, candidate.extracted_text)
        matched, eligible, evidence = match
        if matched < MIN_MATCH_TOKENS:
            continue
        record = (candidate, matched, eligible, evidence)
        if candidate_assignment.family_id == family.id:
            family_matches.append(record)
        else:
            course_matches.append(record)

    if assignment.originality_scope == "family":
        matches = family_matches
    elif assignment.originality_scope == "course":
        matches = family_matches + course_matches
    else:
        matches = family_matches or course_matches
    scope = "assignment_family" if matches and matches[0][0].assignment_id in {
        row[0].assignment_id for row in family_matches
    } else "course"
    receipt = db.get(OriginalityAnalysis, submission.id)
    receipt.candidate_count = (compared_family if assignment.originality_scope == "family" else
                              compared_family + compared_course if assignment.originality_scope == "course" else
                              compared_family if family_matches else compared_family + compared_course)
    receipt.status = "completed" if receipt.candidate_count > 0 else "no_sources"
    reports = []
    for candidate, matched, eligible, evidence in matches:
        report = OriginalityReport(
            submission_id=submission.id,
            candidate_submission_id=candidate.id,
            assignment_family_id=family.id,
            match_scope=scope,
            match_ratio=round(matched / eligible, 6) if eligible else 0,
            eligible_token_count=eligible,
            matched_token_count=matched,
            evidence=evidence,
        )
        db.add(report)
        reports.append(report)
    receipt.report_count = len(reports)
    db.flush()
    return reports


def _compare_submission(db: Session, submission: AcademicSubmission) -> list[OriginalityReport]:
    assignment = db.get(AcademicAssignment, submission.assignment_id)
    existing = db.get(OriginalityAnalysis, submission.id)
    if existing:
        return list(db.scalars(select(OriginalityReport).where(OriginalityReport.submission_id == submission.id)))
    tokens = len(_tokens(submission.extracted_text)[0])
    receipt = OriginalityAnalysis(submission_id=submission.id, status="analyzing",
                                 eligible_token_count=tokens, candidate_count=0,
                                 scope=assignment.originality_scope if assignment else "")
    db.add(receipt)
    db.flush()
    if not assignment or not assignment.originality_enabled:
        receipt.status = "disabled"
    elif tokens < MIN_MATCH_TOKENS:
        receipt.status = "insufficient_text"
    else:
        try:
            with db.begin_nested():
                reports = _compare_submission_pairs(db, submission)
            receipt.completed_at = now_utc()
            return reports
        except Exception:
            logging.getLogger(__name__).error("Deterministic similarity failed for submission %s", submission.id)
            receipt.status = "failed"
    receipt.completed_at = now_utc()
    db.flush()
    return []


def _valid_similarity_report(report, submission, candidate):
    if (not submission or not candidate or submission.student_id == candidate.student_id or
            len(_tokens(candidate.extracted_text)[0]) < MIN_MATCH_TOKENS):
        return False
    eligible, matched, ratio = report.eligible_token_count, report.matched_token_count, report.match_ratio
    if (eligible < MIN_MATCH_TOKENS or eligible != len(_tokens(submission.extracted_text)[0]) or
            not 0 <= matched <= eligible or not math.isfinite(ratio) or
            not 0 <= ratio <= 1 or abs(ratio - matched / eligible) > 0.000001):
        return False
    evidence = report.evidence or []
    if not isinstance(evidence, list):
        return False
    if matched == 0:
        return not evidence
    if not evidence:
        return False
    last_source_end = -1
    for item in evidence:
        if not isinstance(item, dict) or not isinstance(item.get('matched_tokens'), int) or item['matched_tokens'] < MIN_MATCH_TOKENS:
            return False
        for key, text, excerpt in (("source_offsets", submission.extracted_text, "source_excerpt"),
                                   ("candidate_offsets", candidate.extracted_text, "candidate_excerpt")):
            offsets = item.get(key)
            if not isinstance(offsets, dict):
                return False
            start, end = offsets.get("start"), offsets.get("end")
            if (not isinstance(start, int) or not isinstance(end, int) or
                    not 0 <= start < end <= len(text) or text[start:end] != item.get(excerpt)):
                return False
            if len(_tokens(text[start:end])[0]) != item['matched_tokens']:
                return False
            if key == 'source_offsets':
                if start < last_source_end:
                    return False
                last_source_end = end
    return sum(e['matched_tokens'] for e in evidence) == matched


def _report_json(db: Session, report: OriginalityReport) -> dict:
    submission = db.get(AcademicSubmission, report.submission_id)
    candidate = db.get(AcademicSubmission, report.candidate_submission_id)
    assignment = db.get(AcademicAssignment, submission.assignment_id) if submission else None
    candidate_assignment = db.get(AcademicAssignment, candidate.assignment_id) if candidate else None
    offering = db.get(CourseOffering, assignment.offering_id) if assignment else None
    candidate_offering = db.get(CourseOffering, candidate_assignment.offering_id) if candidate_assignment else None
    course = db.get(AcademicCourse, offering.course_id) if offering else None
    candidate_course = db.get(AcademicCourse, candidate_offering.course_id) if candidate_offering else None
    valid = _valid_similarity_report(report, submission, candidate)
    return {
        "analysis_status": "completed" if valid else "incomplete",
        "id": report.id,
        "submission_id": report.submission_id,
        "candidate_submission_id": report.candidate_submission_id,
        "student_id": submission.student_id if submission else None,
        "candidate_student_id": candidate.student_id if candidate else None,
        "assignment_family_id": report.assignment_family_id,
        "match_scope": report.match_scope,
        "match_ratio": report.match_ratio if valid else None,
        "eligible_token_count": report.eligible_token_count,
        "matched_token_count": report.matched_token_count,
        "ratio_definition": "matched normalized tokens / eligible normalized tokens in the submitted text",
        "offset_unit": "Unicode code points, half-open [start, end)",
        "minimum_passage_tokens": MIN_MATCH_TOKENS,
        "evidence": (report.evidence or []) if valid else [],
        "review_status": report.review_status,
        "teacher_notes": report.teacher_notes,
        "created_at": _timestamp(report.created_at),
        "reviewed_at": _timestamp(report.reviewed_at),
        "interpretation": "Similarity evidence for teacher review; not a misconduct finding.",
        "extracted_text": submission.extracted_text if submission else "",
        "document_title": submission.file_name if submission else "",
        "assignment_title": assignment.title if assignment else "",
        "course_name": course.name if course else "",
        "course_code": course.code if course else "",
        "academic_year": offering.academic_year if offering else "",
        "term": offering.term if offering else "",
        "source_title": candidate.file_name if candidate else "",
        "source_assignment": candidate_assignment.title if candidate_assignment else "",
        "source_course": candidate_course.code if candidate_course else "",
        "source_term": (
            f"{candidate_offering.term} {candidate_offering.academic_year}"
            if candidate_offering
            else ""
        ),
    }


@router.on_event("startup")
def seed_academic_demo() -> None:
    from persistence import init_db

    init_db()
    ensure_academic_seed()


@router.get("/courses", response_model=CoursesResponse, summary="[DEMO/LOCAL ONLY] List courses")
def list_courses(
    request: Request,
    role: Literal["student", "teacher"] = Query(...),
) -> dict:
    ensure_academic_seed()
    if auth_required():
        user = request_user(request)
        if (user.role == "student") != (role == "student"):
            raise HTTPException(status_code=403, detail="Requested role does not match credentials.")
        student_id = user.student_id or ""
    else:
        student_id = DEMO_STUDENT_ID
    with SessionLocal() as db:
        courses = db.scalars(select(AcademicCourse).order_by(AcademicCourse.code)).all()
        if auth_required() and role == "teacher":
            courses = [course for course in courses if can_teach(db, user, course.id)]
        items = [_course_dict(db, course, role, student_id) for course in courses]
        return {"items": [item for item in items if role != "student" or item["offerings"]]}


@router.get(
    "/offerings/{offering_id}",
    response_model=OfferingDetailResponse,
    summary="[DEMO/LOCAL ONLY] Get course offering",
)
def get_offering(
    offering_id: str, request: Request, role: Literal["student", "teacher"] = Query("teacher")
) -> dict:
    ensure_academic_seed()
    if auth_required():
        authenticated_role = request_user(request).role
        if (authenticated_role == "student") != (role == "student"):
            raise HTTPException(status_code=403, detail="Requested role does not match credentials.")
        student_id = request_user(request).student_id if role == "student" else None
    else:
        student_id = DEMO_STUDENT_ID if role == "student" else None
    with SessionLocal() as db:
        offering = db.get(CourseOffering, offering_id)
        if offering is None:
            raise HTTPException(status_code=404, detail="Offering not found.")
        if student_id and db.scalar(
            select(CourseEnrollment.id).where(
                CourseEnrollment.offering_id == offering.id,
                CourseEnrollment.student_id == student_id,
            )
        ) is None:
            raise HTTPException(status_code=403, detail="Not enrolled in this offering.")
        course = db.get(AcademicCourse, offering.course_id)
        assignments = db.scalars(
            select(AcademicAssignment)
            .where(AcademicAssignment.offering_id == offering.id)
            .order_by(AcademicAssignment.created_at)
        ).all()
        return {
            "id": offering.id,
            "academic_year": offering.academic_year,
            "term": offering.term,
            "section": offering.section,
            "status": offering.status,
            "class_code": offering.class_code if student_id is None else "",
            "enrollment_count": _active_enrollment_count(db, offering.id) if student_id is None else None,
            "course": {
                "id": course.id,
                "code": (db.get(CourseOwner, course.id).public_code if db.get(CourseOwner, course.id) else course.code),
                "name": course.name,
                "instructor": course.instructor,
                "department": course.department,
            },
            "instructor": course.instructor,
            "assignments": [
                _assignment_view(db, assignment)
                if student_id is None
                else _student_view(_assignment_view(db, assignment))
                for assignment in assignments
            ],
        }


@router.post(
    "/offerings",
    status_code=201,
    response_model=OfferingCreateResponse,
    summary="[DEMO/LOCAL ONLY] Create course offering",
)
def create_offering(payload: OfferingCreate, request: Request) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        course = db.get(AcademicCourse, payload.course_id)
        if course is None:
            raise HTTPException(status_code=404, detail="Course not found.")
        class_code = new_code() if auth_required() and request_user(request).account_id else payload.class_code or new_code()
        if db.scalar(select(CourseOffering.id).where(CourseOffering.class_code == class_code)):
            raise HTTPException(status_code=409, detail="Class code is already in use.")
        offering = CourseOffering(
            course_id=course.id,
            academic_year=payload.academic_year,
            term=payload.term,
            section=payload.section,
            status=payload.status,
            class_code=class_code,
        )
        db.add(offering)
        db.flush()
        if auth_required() and request_user(request).account_id:
            db.add(OfferingPolicy(offering_id=offering.id))
        db.commit()
        db.refresh(offering)
        return {
            "id": offering.id,
            "course_id": offering.course_id,
            "academic_year": offering.academic_year,
            "term": offering.term,
            "section": offering.section,
            "status": offering.status,
            "class_code": offering.class_code,
        }


@router.post(
    "/enroll",
    status_code=201,
    response_model=EnrollmentResponse,
    summary="[DEMO/LOCAL ONLY] Enroll demo student by class code",
)
def enroll(payload: EnrollmentRequest, request: Request) -> dict:
    ensure_academic_seed()
    if auth_required():
        _student_identity(request)
        return join_code(payload, request)
    ensure_academic_seed()
    student_id, student_name = _student_identity(request)
    with SessionLocal() as db:
        offering = db.scalar(
            select(CourseOffering).where(CourseOffering.class_code == payload.class_code.strip())
        )
        if offering is None:
            raise HTTPException(status_code=404, detail="Class code not found.")
        enrollment = db.scalar(
            select(CourseEnrollment).where(
                CourseEnrollment.offering_id == offering.id,
                CourseEnrollment.student_id == student_id,
            )
        )
        created = enrollment is None
        if created:
            enrollment = CourseEnrollment(
                offering_id=offering.id,
                student_id=student_id,
                student_name=student_name,
            )
            db.add(enrollment)
            db.commit()
            db.refresh(enrollment)
        return {
            "id": enrollment.id,
            "offering_id": enrollment.offering_id,
            "student_id": enrollment.student_id,
            "student_name": enrollment.student_name,
            "created": created,
        }


@router.get(
    "/offerings/{offering_id}/assignments",
    response_model=AssignmentItemsResponse,
    summary="[DEMO/LOCAL ONLY] List offering assignments",
)
def list_offering_assignments(
    offering_id: str,
    request: Request,
    role: Literal["student", "teacher"] = Query("teacher"),
) -> dict:
    ensure_academic_seed()
    if auth_required():
        authenticated_role = request_user(request).role
        if (authenticated_role == "student") != (role == "student"):
            raise HTTPException(status_code=403, detail="Requested role does not match credentials.")
        student_id = request_user(request).student_id if role == "student" else None
    else:
        student_id = DEMO_STUDENT_ID if role == "student" else None
    with SessionLocal() as db:
        offering = db.get(CourseOffering, offering_id)
        if offering is None:
            raise HTTPException(status_code=404, detail="Offering not found.")
        if student_id and db.scalar(
            select(CourseEnrollment.id).where(
                CourseEnrollment.offering_id == offering_id,
                CourseEnrollment.student_id == student_id,
            )
        ) is None:
            raise HTTPException(status_code=403, detail="Not enrolled in this offering.")
        assignments = db.scalars(
            select(AcademicAssignment)
            .where(AcademicAssignment.offering_id == offering_id)
            .order_by(AcademicAssignment.created_at)
        ).all()
        return {
            "items": [
                _assignment_view(db, assignment)
                if student_id is None
                else _student_view(_assignment_view(db, assignment))
                for assignment in assignments
            ]
        }


@router.post(
    "/assignments",
    status_code=201,
    response_model=AcademicAssignmentResponse,
    summary="[DEMO/LOCAL ONLY] Create reusable assignment",
)
def create_academic_assignment(payload: AcademicAssignmentCreate) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        offering = db.get(CourseOffering, payload.offering_id)
        if offering is None:
            raise HTTPException(status_code=404, detail="Offering not found.")
        questions = _normalize_questions(payload.questions, payload.maximum_score)
        assignment = _new_academic_assignment(
            db,
            offering,
            payload.title,
            payload.assignment_type,
            payload.description,
            payload.available_at,
            payload.due_at,
            payload.maximum_score,
            payload.allow_resubmission,
            payload.originality_enabled,
            payload.originality_scope,
            questions,
        )
        db.commit()
        db.refresh(assignment)
        return _assignment_view(db, assignment)


@router.post(
    "/assignments/{assignment_id}/reuse",
    status_code=201,
    response_model=AcademicAssignmentResponse,
    summary="[DEMO/LOCAL ONLY] Reuse assignment without submissions or grades",
)
def reuse_assignment(assignment_id: str, payload: AssignmentReuse) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        source = db.get(AcademicAssignment, assignment_id)
        offering = db.get(CourseOffering, payload.offering_id)
        if source is None:
            raise HTTPException(status_code=404, detail="Assignment not found.")
        if offering is None:
            raise HTTPException(status_code=404, detail="Offering not found.")
        if db.get(AcademicCourse, offering.course_id).id != db.get(
            AssignmentFamily, source.family_id
        ).course_id:
            raise HTTPException(status_code=422, detail="Cannot reuse assignment in another course.")
        copied_questions = []
        for question in source.questions or []:
            item = dict(question)
            if not payload.copy_rubric:
                item["criteria"] = []
            if not payload.copy_reference_solution:
                item["reference_solution"] = ""
            if not payload.copy_template:
                item["region"] = None
            copied_questions.append(item)
        if not copied_questions:
            copied_questions = _normalize_questions([], source.maximum_score)
        if payload.copy_template:
            template_regions = list(source.template_regions or [])
            template_rotation = source.template_rotation
        else:
            template_regions = []
            template_rotation = 0
        assignment = _new_academic_assignment(
            db,
            offering,
            source.title,
            source.assignment_type,
            source.description,
            payload.available_at,
            payload.due_at,
            source.maximum_score,
            source.allow_resubmission,
            source.originality_enabled,
            source.originality_scope,
            copied_questions,
            template_regions,
            template_rotation,
            source.family_id,
        )
        db.commit()
        db.refresh(assignment)
        result = _assignment_view(db, assignment)
        result["reused_from"] = source.id
        result["copied"] = {
            "rubric": payload.copy_rubric,
            "reference_solution": payload.copy_reference_solution,
            "template": payload.copy_template,
            "submissions": False,
            "grades": False,
        }
        return result


@router.get(
    "/student/assignments",
    response_model=StudentAssignmentItemsResponse,
    summary="[DEMO/LOCAL ONLY] List demo student assignments",
)
def student_assignments(request: Request) -> dict:
    ensure_academic_seed()
    student_id, _ = _student_identity(request)
    with SessionLocal() as db:
        rows = db.scalars(
            select(AcademicAssignment)
            .join(CourseOffering, AcademicAssignment.offering_id == CourseOffering.id)
            .join(CourseEnrollment, CourseEnrollment.offering_id == CourseOffering.id)
            .where(CourseEnrollment.student_id == student_id)
            .order_by(CourseOffering.academic_year, AcademicAssignment.created_at)
        ).all()
        result = []
        for assignment in rows:
            enrollment = db.scalar(select(CourseEnrollment).where(CourseEnrollment.offering_id == assignment.offering_id, CourseEnrollment.student_id == student_id))
            if not enrollment_active(db, enrollment):
                continue
            item = _student_view(_assignment_view(db, assignment))
            submission = db.scalar(
                select(AcademicSubmission).where(
                    AcademicSubmission.assignment_id == assignment.id,
                    AcademicSubmission.student_id == student_id,
                ).order_by(AcademicSubmission.created_at.desc())
            )
            item["submission"] = (
                {
                    "id": submission.id,
                    "status": submission.status,
                    "submitted_at": _timestamp(submission.created_at),
                    "file_name": submission.file_name,
                }
                if submission
                else None
            )
            grade = db.scalar(
                select(AcademicGrade).where(
                    AcademicGrade.assignment_id == assignment.id,
                    AcademicGrade.submission_id == submission.id if submission else False,
                    AcademicGrade.published.is_(True),
                )
            )
            item["published_grade"] = _published_grade_view(db, grade, assignment) if grade else None
            result.append(item)
        return {"items": result}


@router.get(
    "/student/grades",
    response_model=StudentGradeItemsResponse,
    summary="[DEMO/LOCAL ONLY] List published grades",
)
def student_grades(request: Request) -> dict:
    ensure_academic_seed()
    student_id, _ = _student_identity(request)
    with SessionLocal() as db:
        grades = db.scalars(
            select(AcademicGrade)
            .join(AcademicSubmission, AcademicGrade.submission_id == AcademicSubmission.id)
            .where(
                AcademicSubmission.student_id == student_id,
                AcademicGrade.published.is_(True),
            )
            .order_by(AcademicGrade.published_at.desc())
        ).all()
        items = []
        for grade in grades:
            assignment = db.get(AcademicAssignment, grade.assignment_id)
            submission = db.get(AcademicSubmission, grade.submission_id)
            enrollment = db.scalar(select(CourseEnrollment).where(CourseEnrollment.offering_id == assignment.offering_id, CourseEnrollment.student_id == student_id))
            if not enrollment_active(db, enrollment):
                continue
            items.append(_published_grade_view(db, grade, assignment))
        return {"items": items}


def _published_grade_view(db, grade, assignment):
    return {
        "id": grade.id, "assignment_id": grade.assignment_id,
        "assignment_title": assignment.title, "submission_id": grade.submission_id,
        "points": grade.points, "maximum_score": grade.maximum_score,
        "published_at": _timestamp(grade.published_at),
        **published_feedback(db, grade),
    }


async def _read_submission_file(file: UploadFile, allow_code: bool = False) -> tuple[bytes, str, str]:
    name = Path(file.filename or "").name
    allowed = {"pdf", "docx", "pptx"} | (CODE_SUFFIXES | {"zip"} if allow_code else set())
    if Path(name).suffix.lower().lstrip(".") not in allowed:
        raise HTTPException(415, "Upload a PDF, DOCX or PPTX document.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Submission size must not exceed 30 MB.")
    try:
        structure = await run_in_threadpool(ingest, content, name, allow_code)
    except Exception as exc:
        raise HTTPException(422, str(exc) if isinstance(exc, ValueError) else "Could not read the uploaded document.") from exc
    file._forma_structure = structure
    return content, structure["extracted_text"], structure["file_type"]


@router.post(
    "/assignments/{assignment_id}/submissions",
    status_code=201,
    response_model=SubmissionResponse,
    summary="[DEMO/LOCAL ONLY] Submit PDF or DOCX",
)
async def create_submission(
    assignment_id: str, request: Request, file: UploadFile = File(...), request_id: UUID | None = Form(default=None)
) -> dict:
    ensure_academic_seed()
    student_id, student_name = _student_identity(request)
    if auth_required():
        from accounts import rate_limit
        rate_limit(request, "submission-upload", student_id, 20)
    with SessionLocal() as preflight:
        assignment = preflight.get(AcademicAssignment, assignment_id)
        if assignment is None:
            raise HTTPException(404, "Assignment not found.")
        allow_code = assignment.assignment_type == "programming"
    content, extracted_text, file_type = await _read_submission_file(file, allow_code)
    with SessionLocal() as db:
        assignment = db.get(AcademicAssignment, assignment_id)
        if assignment is None:
            raise HTTPException(status_code=404, detail="Assignment not found.")
        if db.bind.dialect.name == "sqlite":
            db.execute(update(AcademicAssignment).where(AcademicAssignment.id == assignment_id).values(maximum_score=AcademicAssignment.maximum_score))
        db.scalar(select(AcademicAssignment).where(AcademicAssignment.id == assignment_id).with_for_update())
        enrolled = db.scalar(
            select(CourseEnrollment).where(
                CourseEnrollment.offering_id == assignment.offering_id,
                CourseEnrollment.student_id == student_id,
            )
        )
        if not enrollment_active(db, enrolled):
            raise HTTPException(status_code=403, detail="Student is not actively enrolled in this offering.")
        from review_pilot_models import SubmissionUploadReceipt
        content_hash = hashlib.sha256(content).hexdigest()
        if request_id is not None:
            receipt = db.scalar(select(SubmissionUploadReceipt).where(SubmissionUploadReceipt.assignment_id == assignment_id,
                SubmissionUploadReceipt.student_id == student_id, SubmissionUploadReceipt.request_id == str(request_id)))
            if receipt:
                if receipt.content_hash != content_hash or receipt.response['file_name'] != Path(file.filename or '').name[:255]:
                    raise HTTPException(409, 'Upload request ID was already used for a different file.')
                return receipt.response
        now = now_utc()
        available_at = _as_utc(assignment.available_at)
        due_at = _as_utc(assignment.due_at)
        if available_at and now < available_at:
            raise HTTPException(status_code=409, detail="Assignment is not available yet.")
        if due_at and now > due_at:
            raise HTTPException(status_code=409, detail="The submission deadline has passed.")
        prior = db.scalars(
            select(AcademicSubmission).where(
                AcademicSubmission.assignment_id == assignment.id,
                AcademicSubmission.student_id == student_id,
            )
        ).all()
        if prior and not assignment.allow_resubmission:
            raise HTTPException(status_code=409, detail="Resubmission is not allowed.")
        stored_name = f"{uuid4().hex}.{file_type}"
        relative_path = Path(assignment.id) / "academic" / stored_name
        write_blob(relative_path.as_posix(), content)
        legacy_submission_id = None
        if assignment.legacy_assignment_id:
            legacy_path = ""
            page_count = 1
            if file_type == "pdf":
                legacy_path = save_blob(
                    assignment.legacy_assignment_id, "sources", content, ".pdf"
                )
                document = pymupdf.open(stream=content, filetype="pdf")
                page_count = document.page_count
                document.close()
            # Names belong in the authorized academic workspace, never in
            # the legacy anonymous identifier fields. Preserve their separation.
            pseudonym = "Student " + student_id[:8] if auth_required() else student_name
            legacy_submission = Submission(
                assignment_id=assignment.legacy_assignment_id,
                student_name=pseudonym,
                anonymous_student_id=pseudonym,
                source_path=legacy_path,
                page_count=page_count,
                status="ready",
            )
            db.add(legacy_submission)
            db.flush()
            legacy_submission_id = legacy_submission.id
        submission = AcademicSubmission(
            assignment_id=assignment.id,
            student_id=student_id,
            student_name=student_name,
            legacy_submission_id=legacy_submission_id,
            file_path=relative_path.as_posix(),
            file_name=Path(file.filename or stored_name).name[:255],
            file_type=file_type,
            extracted_text=extracted_text,
        )
        db.add(submission)
        db.flush()
        db.add(SubmissionDocument(submission_id=submission.id, content_hash=hashlib.sha256(content).hexdigest(), structure=file._forma_structure))
        reports = _compare_submission(db, submission)
        if auth_required():
            from document_processing import enqueue
            try:
                enqueue(db, submission, assignment, request_user(request).account_id or student_id)
            except ValueError:
                # Receipt is preserved even when today's processing quota is full.
                # No processing success is reported; teacher can request analysis later.
                submission.status = "analysis_not_scheduled"
        result = {
            "id": submission.id,
            "assignment_id": submission.assignment_id,
            "student_id": submission.student_id,
            "file_name": submission.file_name,
            "file_type": submission.file_type,
            "status": submission.status,
            "extracted_text": submission.extracted_text,
            "submitted_at": _timestamp(submission.created_at),
            "originality_reports": (
                []
                if auth_required()
                else [_report_json(db, report) for report in reports]
            ),
        }
        if request_id is not None:
            db.add(SubmissionUploadReceipt(assignment_id=assignment_id, student_id=student_id, request_id=str(request_id),
                                           content_hash=content_hash, response=result))
        db.commit()
        return result


@router.get(
    "/teacher/originality",
    response_model=OriginalityItemsResponse,
    summary="[DEMO/LOCAL ONLY] List originality reports for teacher review",
)
def teacher_originality(
    request: Request,
    offering_id: str | None = None,
    review_status: Literal["unreviewed", "reviewed", "follow_up"] | None = None,
) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        query = select(OriginalityReport).order_by(OriginalityReport.created_at.desc())
        if offering_id:
            query = (
                query.join(
                    AcademicSubmission,
                    OriginalityReport.submission_id == AcademicSubmission.id,
                )
                .join(
                    AcademicAssignment,
                    AcademicSubmission.assignment_id == AcademicAssignment.id,
                )
                .where(AcademicAssignment.offering_id == offering_id)
            )
        if review_status:
            query = query.where(OriginalityReport.review_status == review_status)
        reports = db.scalars(query).all()
        if auth_required():
            user = request_user(request)
            reports = [report for report in reports if can_teach(db, user, db.get(CourseOffering, db.get(AcademicAssignment, db.get(AcademicSubmission, report.submission_id).assignment_id).offering_id).course_id)]
        return {"items": [_report_json(db, report) for report in reports]}


@router.get(
    "/submissions/{submission_id}/originality",
    response_model=SubmissionOriginalityResponse,
    summary="[DEMO/LOCAL ONLY] Get submission similarity evidence",
)
def submission_originality(submission_id: str) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        submission = db.get(AcademicSubmission, submission_id)
        if submission is None:
            raise HTTPException(status_code=404, detail="Submission not found.")
        reports = db.scalars(
            select(OriginalityReport)
            .where(OriginalityReport.submission_id == submission_id)
            .order_by(OriginalityReport.match_ratio.desc())
        ).all()
        items = [_report_json(db, r) for r in reports]
        receipt = db.get(OriginalityAnalysis, submission_id)
        from similarity_status import similarity_summary
        summary = similarity_summary(submission, reports, receipt,
            {r.candidate_submission_id: db.get(AcademicSubmission, r.candidate_submission_id) for r in reports})
        status, eligible = summary['analysis_status'], summary['eligible_token_count']
        assignment = db.get(AcademicAssignment, submission.assignment_id)
        offering = db.get(CourseOffering, assignment.offering_id)
        course = db.get(AcademicCourse, offering.course_id)
        return {"submission_id": submission_id, "items": items, "analysis": {
            "analysis_status": status, "submission_id": submission_id, "document_title": submission.file_name,
            "assignment_title": assignment.title, "course_name": course.name, "academic_year": offering.academic_year,
            "extracted_text": submission.extracted_text, "eligible_token_count": eligible,
            "candidate_count": receipt.candidate_count if receipt else None,
            "match_ratio": max((r.match_ratio for r in reports), default=0) if status == "completed" else None,
            "match_scope": receipt.scope if receipt else (reports[0].match_scope if reports else assignment.originality_scope),
            "created_at": _timestamp(receipt.completed_at) if receipt else (_timestamp(reports[0].created_at) if reports else None),
        }}


@router.patch(
    "/reports/{report_id}/review",
    response_model=OriginalityReportResponse,
    summary="[DEMO/LOCAL ONLY] Record teacher review, not a misconduct conclusion",
)
def review_originality_report(report_id: str, payload: ReviewUpdate) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        report = db.get(OriginalityReport, report_id)
        if report is None:
            raise HTTPException(status_code=404, detail="Originality report not found.")
        if not _valid_similarity_report(report, db.get(AcademicSubmission, report.submission_id),
                                        db.get(AcademicSubmission, report.candidate_submission_id)):
            raise HTTPException(409, "Similarity evidence is incomplete; a review cannot be attached.")
        report.review_status = payload.review_status
        report.teacher_notes = payload.teacher_notes
        report.reviewed_at = now_utc()
        db.commit()
        db.refresh(report)
        return _report_json(db, report)


@router.get(
    "/teacher/assignments/{assignment_id}/archive",
    response_model=ArchiveItemsResponse,
    summary="[DEMO/LOCAL ONLY] View historical assignment submissions",
)
def assignment_archive(assignment_id: str, request: Request) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        assignment = db.get(AcademicAssignment, assignment_id)
        if assignment is None:
            raise HTTPException(status_code=404, detail="Assignment not found.")
        family_assignments = db.scalars(
            select(AcademicAssignment).where(
                AcademicAssignment.family_id == assignment.family_id
            )
        ).all()
        results = []
        for family_assignment in family_assignments:
            offering = db.get(CourseOffering, family_assignment.offering_id)
            if auth_required() and not can_teach(db, request_user(request), offering.course_id):
                continue
            for submission in db.scalars(
                select(AcademicSubmission).where(
                    AcademicSubmission.assignment_id == family_assignment.id
                )
            ).all():
                results.append(
                    {
                        "submission_id": submission.id,
                        "assignment_id": family_assignment.id,
                        "assignment_title": family_assignment.title,
                        "academic_year": offering.academic_year,
                        "term": offering.term,
                        "student_id": submission.student_id,
                        "student_name": student_display_name(db, submission.student_id, submission.student_name),
                        "file_name": submission.file_name,
                        "file_type": submission.file_type,
                        "submitted_at": _timestamp(submission.created_at),
                        "extracted_text": submission.extracted_text,
                    }
                )
        return {"items": sorted(results, key=lambda item: (item["academic_year"], item["submitted_at"]))}


@router.post(
    "/assignments/{assignment_id}/publish",
    response_model=PublishResponse,
    summary="[DEMO/LOCAL ONLY] Publish approved legacy Answer grades to enrolled students",
)
def publish_grades(assignment_id: str) -> dict:
    ensure_academic_seed()
    with SessionLocal() as db:
        assignment = db.get(AcademicAssignment, assignment_id)
        if assignment is None:
            raise HTTPException(status_code=404, detail="Assignment not found.")
        if assignment.legacy_assignment_id is None:
            raise HTTPException(status_code=409, detail="Assignment has no linked grading record.")
        eligible_submissions = db.scalars(
            select(AcademicSubmission).where(
                AcademicSubmission.assignment_id == assignment.id,
                AcademicSubmission.legacy_submission_id.is_not(None),
            )
        ).all()
        published = []
        for academic_submission in eligible_submissions:
            answers = db.scalars(
                select(Answer).where(
                    Answer.assignment_id == assignment.legacy_assignment_id,
                    Answer.submission_id == academic_submission.legacy_submission_id,
                    Answer.approved.is_(True),
                )
            ).all()
            if not answers:
                continue
            points = sum(
                answer.human_score
                if answer.human_score is not None
                else (answer.suggested_score or 0)
                for answer in answers
            )
            grade = db.scalar(
                select(AcademicGrade).where(
                    AcademicGrade.assignment_id == assignment.id,
                    AcademicGrade.submission_id == academic_submission.id,
                )
            )
            if grade is not None and db.get(PublishedDocumentReview, grade.id):
                continue  # Universal review remains authoritative until explicitly republished.
            if points > assignment.maximum_score:
                raise HTTPException(409, "Approved answer points exceed the assignment maximum.")
            if grade is None:
                grade = AcademicGrade(
                    assignment_id=assignment.id,
                    submission_id=academic_submission.id,
                    points=points,
                    maximum_score=assignment.maximum_score,
                )
                db.add(grade)
            grade.points = points
            grade.maximum_score = assignment.maximum_score
            grade.published = True
            grade.published_at = now_utc()
            published.append(
                {
                    "submission_id": academic_submission.id,
                    "points": points,
                    "maximum_score": assignment.maximum_score,
                    "approved_answer_count": len(answers),
                }
            )
        db.commit()
        return {
            "assignment_id": assignment.id,
            "published_count": len(published),
            "grades": published,
            "notice": "Only approved legacy Answer grades were published.",
        }


__all__ = ["router", "ensure_academic_seed"]
