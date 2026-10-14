import io
import zipfile
import pytest
import pymupdf
from docx import Document
from pptx import Presentation
from services.document_ingestion import ingest, MAX_BYTES


def pdf(pages):
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((48, 72), text)
    result = document.tobytes(); document.close(); return result


def test_multipage_pdf_preserves_pages_and_reports_scans_without_inventing_answers():
    structure = ingest(pdf(["Question 1\nWork begins here", "continued on next page", ""]), "answers.pdf")
    assert structure["unit_count"] == 3
    assert structure["pages"][1]["number"] == 2
    assert structure["pages"][0]["regions"]
    assert structure["question_associations"] == []
    assert "no embedded text" in structure["warnings"][0]
    assert structure["requires_human_review"]


def test_docx_tables_no_fabricated_page_numbers():
    document = Document(); document.add_heading("Introduction", 1); document.add_paragraph("A synthetic laboratory report.")
    table = document.add_table(rows=2, cols=2); table.cell(0, 0).text = "Measurement"; table.cell(1, 1).text = "12.3"
    output = io.BytesIO(); document.save(output)
    result = ingest(output.getvalue(), "report.docx")
    assert result["pages"][0]["unit"] == "document"
    assert "12.3" in result["extracted_text"]
    assert result["pages"][0]["tables"]
    assert "page numbers" in result["warnings"][0]


def test_pptx_slide_text_and_explicit_visual_limitations():
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "Entropy"; slide.placeholders[1].text = "Synthetic lecture notes."
    output = io.BytesIO(); presentation.save(output)
    result = ingest(output.getvalue(), "slides.pptx")
    assert result["pages"][0]["number"] == 1 and result["pages"][0]["unit"] == "slide"
    assert "Entropy" in result["extracted_text"]
    assert "rendering" in result["warnings"][0]


@pytest.mark.parametrize("content,filename", [(b"not pdf", "x.pdf"), (b"PKbad", "x.docx"), (b"MZexecutable", "x.pptx"), (b"", "x.pdf"), (b"bad", "x.exe")])
def test_malformed_and_unsupported_documents(content, filename):
    with pytest.raises(Exception): ingest(content, filename)


def test_limits_and_ooxml_path_traversal_macros():
    with pytest.raises(ValueError, match="1–50"): ingest(pdf([""] * 51), "x.pdf")
    with pytest.raises(ValueError, match="30 MB"): ingest(b"x" * (MAX_BYTES + 1), "x.pdf")
    for malicious in ["../escape.xml", "word/vbaProject.bin", "word/embeddings/object.bin"]:
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>"); archive.writestr("word/document.xml", "<document/>"); archive.writestr(malicious, "bad")
        with pytest.raises(ValueError): ingest(output.getvalue(), "x.docx")


def test_code_is_static_and_archives_restricted():
    result = ingest(b"print('synthetic')", "answer.py", allow_code=True)
    assert result["pages"][0]["unit"] == "file"
    assert "never executed" in result["warnings"][0]
    with pytest.raises(ValueError): ingest(b"print('x')", "answer.py")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive: archive.writestr("../../attack.py", "bad")
    with pytest.raises(ValueError): ingest(output.getvalue(), "answer.zip", allow_code=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive: archive.writestr("answer.py", "print('x')")
    assert ingest(output.getvalue(), "answer.zip", allow_code=True)["unit_count"] == 1
